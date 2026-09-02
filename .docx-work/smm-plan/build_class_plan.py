from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(r"C:\Users\LENOVO\OneDrive\Documents\claude cowork\architech\smm\SMM_Course_Syllabus_and_Class_Plan.docx")
OUTPUT = Path(r"C:\Users\LENOVO\OneDrive\Documents\programming\disegno\SMM_Applied_Class_Plan.docx")

NAVY = "365F91"
MID_BLUE = "D9EAF7"
PALE_BLUE = "EDF4FA"
PALE_GRAY = "F2F2F2"
DARK = RGBColor(31, 31, 31)
MUTED = RGBColor(89, 89, 89)
BLUE = RGBColor(54, 95, 145)
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, size=None, bold=None, color=None, italic=None, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=0):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            cell.width = Inches(width / 1440)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_body_text(paragraph, text, size=8.8, bold=False, color=DARK, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    return run


def add_bullet(doc, text, size=8.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    set_keep_with_next(p, False)
    add_body_text(p, text, size=size)
    return p


def add_heading(doc, text, before=4, after=2):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=11.2, bold=True, color=BLUE)
    return p


def set_cell_text(cell, text, size=8.0, bold=False, color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def remove_body_content(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def update_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = "01-Sep-26\t\tSMM_Applied_Class_Plan"
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, size=8, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(8.8)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Aptos Display"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    h1.font.size = Pt(11.2)
    h1.font.bold = True
    h1.font.color.rgb = BLUE
    h1.paragraph_format.keep_with_next = True


def add_overview_table(doc):
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [1350, 3330, 1350, 3330])
    entries = [
        ("Schedule", "Saturday & Sunday, 1:00-2:00 pm"),
        ("Duration", "10 weeks / 20 live sessions"),
        ("Core method", "Real business work + weekly experiments"),
        ("Teacher role", "Mentor/manager; remote support during the week"),
    ]
    for idx, (label, value) in enumerate(entries):
        row = idx // 2
        col = (idx % 2) * 2
        set_cell_text(table.cell(row, col), label, size=8.2, bold=True, color=WHITE)
        shade_cell(table.cell(row, col), NAVY)
        set_cell_text(table.cell(row, col + 1), value, size=8.2)
        shade_cell(table.cell(row, col + 1), PALE_BLUE)
    return table


def add_weekly_rhythm_table(doc):
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [3120, 3120, 3120])
    headers = ["SATURDAY | REVIEW & DEFEND", "SUNDAY | BUILD & COMMIT", "MON-FRI | FIELDWORK"]
    body = [
        "Progress evidence; student presentations; questions and defense; manager-style feedback.",
        "Short concept input; guided planning; task assignment; student states the week's commitment.",
        "Students execute with the business, record evidence and metrics, and request remote help when blocked.",
    ]
    for col in range(3):
        set_cell_text(table.cell(0, col), headers[col], size=8.0, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(table.cell(0, col), NAVY)
        set_cell_text(table.cell(1, col), body[col], size=8.0)
        shade_cell(table.cell(1, col), PALE_BLUE if col != 1 else PALE_GRAY)
    set_repeat_table_header(table.rows[0])
    return table


def build():
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    remove_body_content(doc)
    configure_styles(doc)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    update_footer(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("Applied Social Media Marketing Class Plan"), size=15.5, bold=True, color=BLUE, name="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(4)
    subtitle.paragraph_format.keep_with_next = True
    set_run_font(subtitle.add_run("Real local businesses | Experiment-driven learning | 10 weekends | 20 x 1-hour classes"), size=8.7, color=MUTED, italic=True)

    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.10)
    lead.paragraph_format.right_indent = Inches(0.10)
    lead.paragraph_format.space_before = Pt(0)
    lead.paragraph_format.space_after = Pt(4)
    set_paragraph_shading(lead, MID_BLUE)
    add_body_text(
        lead,
        "Purpose: Students learn the basics of branding and marketing by managing a scoped, approved social media project for an assigned local business. Each week they plan, execute, measure, present, and improve real work.",
        size=9.0,
        bold=True,
    )

    add_overview_table(doc)

    add_heading(doc, "How the course works", before=4, after=1)
    add_bullet(doc, "Each student is assigned one local business and begins with a short discovery interview, approval boundaries, and baseline audit.")
    add_bullet(doc, "Learning follows a weekly cycle: observe a business need, form a hypothesis, create and execute, collect evidence, explain the result, and decide the next action.")
    add_bullet(doc, "Students do the work independently. The teacher gives concise inputs, sets standards, removes blockers, and manages deadlines rather than delivering long lectures.")
    add_bullet(doc, "Nothing is published, messaged, advertised, or paid for without the business owner's approval. Student work must respect privacy, copyright, brand voice, and platform rules.")

    add_heading(doc, "Weekend-to-weekday operating rhythm", before=3, after=2)
    add_weekly_rhythm_table(doc)

    add_heading(doc, "Learning outcomes", before=3, after=1)
    for text in [
        "Explain a brand's audience, positioning, promise, voice, and visual direction in plain language.",
        "Connect marketing goals to a simple customer journey, offer, channel, content, call to action, and metric.",
        "Create basic social content and run at least two small, measurable experiments for the assigned business.",
        "Use evidence to defend choices, answer questions, accept critique, and recommend a clear next step.",
    ]:
        add_bullet(doc, text, size=8.3)

    add_heading(doc, "Evidence and assessment", before=3, after=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    pieces = [
        ("50%", " weekly business deliverables and experiment evidence  |  "),
        ("25%", " presentations, defense, and clarity  |  "),
        ("15%", " reliability, client communication, and deadline management  |  "),
        ("10%", " final portfolio and business handoff"),
    ]
    for score, desc in pieces:
        set_run_font(p.add_run(score), size=8.3, bold=True, color=BLUE)
        set_run_font(p.add_run(desc), size=8.3, color=DARK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.keep_together = True
    add_body_text(note, "Weekly proof: task tracker, links/screenshots or drafts, business feedback, metrics, and a short reflection: What happened? Why? What next?", size=8.1, italic=True, color=MUTED)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    page2_title = doc.add_paragraph()
    page2_title.paragraph_format.space_before = Pt(0)
    page2_title.paragraph_format.space_after = Pt(3)
    page2_title.paragraph_format.keep_with_next = True
    set_run_font(page2_title.add_run("10-Week / 20-Session Execution Roadmap"), size=13.5, bold=True, color=BLUE, name="Aptos Display")

    roadmap = [
        ("1", "Business onboarding & baseline", "Course model; meet the business; define scope and approval rules.", "Discovery questions; audit channels, audience, offer, and current performance.", "Complete interview, consent/scope sheet, baseline screenshots and 3 priority problems."),
        ("2", "Branding basics", "Present baseline; defend the chosen priority problem.", "Audience, positioning, brand promise, voice, visual consistency.", "Create a one-page brand brief with audience, promise, voice, colors/type, and examples."),
        ("3", "Marketing basics", "Present brand brief; answer client-fit questions.", "Customer journey, objectives, offer, channel, CTA, and useful metrics.", "Map a simple funnel; write 2 measurable experiment hypotheses and success criteria."),
        ("4", "Content planning", "Pitch hypotheses and receive challenge questions.", "Content pillars, formats, calendar, copy structure, and approval workflow.", "Build a 1-week calendar and produce 3 approval-ready content drafts."),
        ("5", "Experiment 1: execute", "Content critique; defend creative and CTA choices.", "Finalize organic experiment; publishing checklist; evidence capture.", "Launch approved Experiment 1; record dates, content, reach, engagement, clicks/leads."),
        ("6", "Audience communication", "Report early signals; explain what is and is not working.", "Comments/DMs, response tone, community care, escalation, and privacy.", "Use an approved response guide; document interactions and recurring questions."),
        ("7", "Analytics & iteration", "Present Experiment 1 results and defend the conclusion.", "Read core metrics; compare baseline; separate signal from assumption.", "Write a 1-page experiment review and recommend one evidence-based change."),
        ("8", "Experiment 2: improve", "Pitch the change and expected result.", "Design a controlled improvement: hook, format, timing, offer, CTA, or audience.", "Launch approved Experiment 2; keep other variables stable where practical; log results."),
        ("9", "Micro-campaign & reporting", "Compare experiments; explain trade-offs and limits.", "Combine brand, content, CTA, and measurement into a short campaign plan.", "Run/prepare an approved 1-week micro-campaign and draft a client-facing report."),
        ("10", "Final defense & handoff", "Final presentation: problem, actions, evidence, learning, recommendation.", "Peer/client questions; revise portfolio; plan a sustainable next 30 days.", "Deliver final portfolio, asset folder, metrics summary, recommendations, and client handoff."),
    ]

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    widths = [520, 1500, 2200, 2200, 2940]
    set_table_geometry(table, widths)
    headers = ["WK", "FOCUS", "SATURDAY 1-2 PM", "SUNDAY 1-2 PM", "MONDAY-FRIDAY DELIVERABLE"]
    for col, header in enumerate(headers):
        set_cell_text(table.cell(0, col), header, size=7.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(table.cell(0, col), NAVY)
    set_repeat_table_header(table.rows[0])

    for row_index, values in enumerate(roadmap, start=1):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[col], value, size=7.25 if col != 1 else 7.4, bold=(col in (0, 1)), align=align)
            if row_index % 2 == 0:
                shade_cell(cells[col], PALE_BLUE)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for cell in cells[2:]:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    set_table_geometry(table, widths)

    completion = doc.add_paragraph()
    completion.paragraph_format.space_before = Pt(3)
    completion.paragraph_format.space_after = Pt(0)
    completion.paragraph_format.keep_together = True
    set_paragraph_shading(completion, PALE_GRAY)
    add_body_text(completion, "Completion standard: Every student must show approved real-business work, evidence from two experiments, reflective decisions, a clear verbal defense, and a usable handoff for the business.", size=8.1, bold=True)

    doc.core_properties.title = "Applied Social Media Marketing Class Plan"
    doc.core_properties.subject = "10-week experiment-driven class plan for assigned local businesses"
    doc.core_properties.keywords = "SMM, branding, marketing, local business, experiment-driven learning"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
